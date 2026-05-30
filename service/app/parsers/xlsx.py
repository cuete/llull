"""XLSX parser using openpyxl."""
from __future__ import annotations

from pathlib import Path

import structlog

from app.parsers.base import BaseParser

log = structlog.get_logger()


class XlsxParser(BaseParser):
    """Extract text from Excel files."""

    async def extract(self, source: Path | str) -> str:
        try:
            import openpyxl
        except ImportError as e:
            raise ImportError("openpyxl is required for XLSX parsing") from e

        path = Path(source) if isinstance(source, str) else source
        if not path.exists():
            raise FileNotFoundError(f"XLSX file not found: {path}")

        log.info("xlsx_extract_start", path=str(path))
        rows: list[str] = []

        try:
            wb = openpyxl.load_workbook(str(path), data_only=True)
            for sheet_name in wb.sheetnames:
                sheet = wb[sheet_name]
                rows.append(f"## Sheet: {sheet_name}")
                for row in sheet.iter_rows(values_only=True):
                    row_text = "\t".join(str(cell) if cell is not None else "" for cell in row)
                    if row_text.strip():
                        rows.append(row_text)
        except Exception as e:
            log.error("xlsx_extract_failed", path=str(path), error=str(e))
            raise

        extracted = "\n".join(rows)
        log.info("xlsx_extract_done", path=str(path), chars=len(extracted))
        return extracted
